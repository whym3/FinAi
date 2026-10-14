from datetime import date

import pytest
from openpyxl import Workbook

from app.imports.kinds import auto_map, detect, missing_required
from app.imports.readers import cell_date, header_and_preview, split_code_name
from app.imports.run import write_rows
from app.models import CustomerInvoice, Party, PurchaseOrder

PR_HEADERS = ['Plant', 'Document Date', 'Purchasing Document', 'Item', 'Name of Supplier', 'Material', 'Material Group', 'Short Text', 'Order Quantity', 'Order Unit',
              'Still to be delivered (qty)', 'Net Price', 'Net Order Value', 'Release State', 'Purchasing Doc. Type']
SR_HEADERS = ['Marketing Person', 'Profit Center', 'Invoice Date', 'GST Invoice Number', 'Sales Order No.', 'Invoice Number', 'Transaction Type', 'Bill to Party Code', 'Bill to Party Name',
              'Bill to Party State', 'Bill to Party City', 'State Zone', 'Material Group', 'Sub Material group', 'Material Number', 'Material Description', 'Total GST', 'Quantity In TON',
              'Taxable Value in INR', 'Total Tax Value', 'Invoice Value in INR', 'Quantity In PC', 'Recon Account Desc.', 'Tax Code', 'Vehicle Number', 'GR No']


def test_detect_registers():
    kind, score, m = detect(PR_HEADERS)
    assert kind == 'purchase_register' and m['po_no'] == 2 and m['open_qty'] == 10 and m['vendor'] == 4 and not missing_required(kind, m)
    kind, score, m = detect(SR_HEADERS)
    assert kind == 'sales_register' and m['invoice_no'] == 5 and m['gst_invoice_no'] == 3 and m['taxable'] == 18 and m['customer_code'] == 7 and not missing_required(kind, m)
    assert detect(['Foo', 'Bar', 'Baz'])[0] is None


def test_parsers():
    assert split_code_name('1005750    BLEND COLOURS P') == ('1005750', 'BLEND COLOURS P')
    assert split_code_name('Deccan Steel') == (None, 'Deccan Steel')
    assert cell_date(46266) == date(2026, 9, 1) and cell_date('01-09-2026') == date(2026, 9, 1) and cell_date(46266.0) == date(2026, 9, 1)


@pytest.fixture
def xlsx(tmp_path):
    def make(name, headers, rows, pivot_first=False):
        wb = Workbook()
        if pivot_first:
            ws0 = wb.active
            ws0.title = 'Sheet2'
            ws0.append(['Row Labels', 'Sum'])
            ws0.append(['A', 1])
            ws = wb.create_sheet('Sheet1')
        else:
            ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        p = tmp_path / name
        wb.save(p)
        return str(p)
    return make


def test_header_detection_and_purchase_register_import(db, firm, xlsx):
    f, e, admin = firm
    rows = [['1200', 46295, '4110008142', 10, '1005750    BLEND COLOURS P', 'RMUV', 'WATERTANK', 'UV STABILIZER', 600, 'KG', 100, 207, 124200, 'XXX', 'ZRAW'],
            ['1900', 46295, '4110008143', 10, '1006192    JAIMARUTI POLYC', 'RMRES', 'RMRESU', 'PVC RESIN', 116, 'MT', 116, 97500, 11310000, 'XXX', 'ZRAW'],
            ['1900', 46295, '4110008143', 20, '1006192    JAIMARUTI POLYC', 'RMRES2', 'RMRESU', 'PVC RESIN 2', 10, 'MT', 0, 1000, 10000, 'XXX', 'ZRAW'],
            ['1900', 46295, '', 30, 'bad row', '', '', 'no po', 1, 'MT', 0, 1, 1, '', 'ZRAW']]
    path = xlsx('pr.xlsx', PR_HEADERS, rows)
    h, headers, preview, total = header_and_preview(path, 'pr.xlsx', 'Sheet')
    assert h == 0 and headers == PR_HEADERS and total == 4
    kind, _, mapping = detect(headers)

    def raw_rows():
        from app.imports.readers import iter_rows
        for i, r in enumerate(iter_rows(path, 'pr.xlsx', 'Sheet')):
            if i > h:
                yield r
    res = write_rows(db, e.id, kind, raw_rows(), None, mapping)
    assert res['rows'] == 4 and res['skipped'] == 1 and res['created'] == 2 and res['errors'][0]['message'].startswith('PO number')
    assert res['period_from'] == date(2026, 9, 30)
    po = db.query(PurchaseOrder).filter_by(entity_id=e.id, po_no='4110008143').one()
    assert len(po.lines) == 2 and float(po.lines[0].open_qty) == 116 and po.plant == '1900' and po.vendor_name == 'JAIMARUTI POLYC'
    v = db.query(Party).filter_by(entity_id=e.id, code='1006192').one()
    assert v.kind == 'vendor' and po.vendor_id == v.id
    # re-import = update, not duplicate
    res2 = write_rows(db, e.id, kind, raw_rows(), None, mapping)
    assert res2['created'] == 0 and res2['updated'] == 2 and db.query(PurchaseOrder).filter_by(entity_id=e.id).count() == 2
    assert {p.code for p in e.plants} == {'1200', '1900'}


def test_sales_register_import_groups_lines(db, firm, xlsx):
    f, e, admin = firm
    base = ['Vivek', '0000001700', 46266, 'RP2627000229', '0000098639', '0090136208', 'ZF2', '2005154', 'M/S PRINCE MARKETING', 'Chhattisgarh', 'RAIPUR', 'RAIPUR', 'SWR FIT', 'SWR', 'PMN', 'TRAP', 18, 0]
    rows = [base + [3031.71, 545.70, 3577.41, 36, 'SUNDRY DEBTORS', 'AC', 'CG04JD8920', '1700SRS'],
            base + [6346.08, 1142.30, 7488.38, 144, 'SUNDRY DEBTORS', 'AC', 'CG04JD8920', '1700SRS'],
            ['Vivek', '0000001700', 46270, 'RP2627000300', '0000098700', '0090136300', 'ZRE', '2005154', 'M/S PRINCE MARKETING', 'Chhattisgarh', 'RAIPUR', 'RAIPUR', 'SWR FIT', 'SWR', 'PMN', 'TRAP', 18, 0, 100, 18, 118, 1, 'SUNDRY DEBTORS', 'AC', '', '']]
    path = xlsx('sr.xlsx', SR_HEADERS, rows, pivot_first=True)
    from app.imports.readers import iter_rows, sheet_names
    assert sheet_names(path, 'sr.xlsx') == ['Sheet2', 'Sheet1']
    kind, _, mapping = detect(SR_HEADERS)
    res = write_rows(db, e.id, kind, (r for i, r in enumerate(iter_rows(path, 'sr.xlsx', 'Sheet1')) if i > 0), None, mapping)
    assert res['created'] == 2 and res['skipped'] == 0
    inv = db.query(CustomerInvoice).filter_by(entity_id=e.id, invoice_no='0090136208').one()
    assert len(inv.lines) == 2 and round(float(inv.total), 2) == 11065.79 and inv.date == date(2026, 9, 1) and inv.due_date == date(2026, 10, 1) and inv.gst_invoice_no == 'RP2627000229'
    cm = db.query(CustomerInvoice).filter_by(entity_id=e.id, invoice_no='0090136300').one()
    assert cm.doc_type == 'credit_memo'
    c = db.query(Party).filter_by(entity_id=e.id, code='2005154').one()
    assert c.kind == 'customer' and c.state == 'Chhattisgarh' and inv.customer_id == c.id
